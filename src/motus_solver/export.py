from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet


def open_or_create_workbook(path: str | Path) -> Workbook:
    """Ouvre le classeur multi-sheet existant à `path`, ou en crée un nouveau.

    Point d'entrée commun pour tous les exports du solveur : chaque export ajoute
    son propre sheet via `add_sheet` sur ce même classeur plutôt que de générer un
    fichier séparé, pour rester cohérent avec un unique classeur multi-sheet.
    """
    path = Path(path)
    if path.exists():
        return load_workbook(path)
    workbook = Workbook()
    workbook.remove(workbook.active)  # pas de sheet "Sheet" par défaut sans nom
    return workbook


def add_sheet(
    workbook: Workbook, name: str, headers: list[str], rows: list[list[Any]]
) -> Worksheet:
    """Ajoute (ou remplace si déjà présent) un sheet nommé `name` : une ligne d'en-tête
    puis une ligne par élément de `rows`."""
    if name in workbook.sheetnames:
        del workbook[name]
    worksheet = workbook.create_sheet(title=name)
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    return worksheet


def save_workbook(workbook: Workbook, path: str | Path) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
